"""Excel export: packing list in the factory's 'Output Want' layout + warehouse instructions."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
MIX_FILL = PatternFill("solid", fgColor="FFF4D6")
LOW_FILL = PatternFill("solid", fgColor="FADBD8")


def _hdr(ws, row, headers, widths=None):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def _sku_map(res):
    return {s["idx"]: s for s in res["skus"]}


def _orient_text(cont):
    return " | ".join(o["text"] for o in cont["orientations"])


def build_workbook(res: dict) -> bytes:
    wb = Workbook()
    meta = res.get("meta", {}) or {}
    sm = res["summary"]

    # ---------------------------------------------------------------- Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "KWIK PATCH PRIVATE LIMITED - Carton Packing Plan"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = meta.get("title", "")
    ws["A3"] = meta.get("consignee", "")
    rows = [
        ("Full cartons", sm["full_cartons"]),
        ("Mixed (leftover) cartons", sm["mixed_cartons"]),
        ("TOTAL cartons", sm["total_cartons"]),
        ("Inner boxes in mixed cartons", sm["mixed_boxes"]),
        ("Units in mixed cartons", sm["mixed_units"]),
        ("Average fill of mixed cartons", f"{(sm['avg_mixed_fill'] or 0) * 100:.1f} %"),
        ("Theoretical minimum mixed cartons (volume bound)", sm["mixed_lower_bound"]),
        ("Orientation rule", "All boxes in a carton in ONE direction" if res["settings"]["mode"] == "uniform"
         else "One direction per product inside a carton"),
        ("Cartons allowed for mixed packing", ", ".join(res.get("allowed_cartons", []))),
        ("Carton calculation size", f"Outer carton dimensions reduced by {res['settings'].get('carton_reduction_inch', 1.0):g} inch (L, W, H) for internal usable space"),
    ]
    for i, (k, v) in enumerate(rows, start=5):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 60

    # ---------------------------------------------------------------- Packing list (same layout as 'Output Want')
    ws = wb.create_sheet("Packing List")
    _hdr(ws, 1, ["PRODUCT Code", "DESCRIPTION OF GOODS", "CARTON NOS", "NO OF CARTONS\n(units/box X boxes X cartons)",
                 "TOTAL QUANTITY IN UNITS", "DIMENSION OF THE CARTONS (IN INCHES)", "TYPE"],
         [16, 48, 14, 24, 18, 28, 10])
    r = 2
    for row in res["packing_list"]:
        vals = [row["code"], row["name"], row["carton_nos"], row["formula"], row["units"], row["dimension"], row["type"]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BORDER
            if row["type"] == "MIXED":
                cell.fill = MIX_FILL
        r += 1
    ws.cell(row=r + 1, column=2, value="TOTAL CARTONS").font = Font(bold=True)
    ws.cell(row=r + 1, column=3, value=sm["total_cartons"]).font = Font(bold=True)
    ws.freeze_panes = "A2"

    # ---------------------------------------------------------------- Mixed cartons - contents
    ws = wb.create_sheet("Mixed Cartons")
    _hdr(ws, 1, ["Carton No", "Carton", "Carton size (in)", "Fill %", "Product code", "Product", "Inner boxes",
                 "Units", "Inner box size L x W x H (mm)", "How to place the box (direction)"],
         [10, 10, 20, 8, 14, 42, 11, 10, 24, 80])
    r = 2
    for m in res["mixed"]:
        first = True
        for cont in m["contents"]:
            vals = [m["no"], m["carton"]["name"], m["carton"]["label"] if first else "",
                    round(m["fill"] * 100, 1) if first else "", cont["code"], cont["name"], cont["boxes"],
                    cont["units"], " x ".join(f"{d:g}" for d in cont["dims"]), _orient_text(cont)]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=r, column=c, value=v)
                cell.border = BORDER
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if m.get("low_fill"):
                    cell.fill = LOW_FILL
            first = False
            r += 1
    ws.freeze_panes = "A2"

    # ---------------------------------------------------------------- Layer instructions
    ws = wb.create_sheet("Layer Instructions")
    _hdr(ws, 1, ["Carton No", "Carton", "Layer", "Step", "Action", "Product", "Boxes", "Grid (along length x along width)",
                 "Start: from left wall (mm)", "Start: from front wall (mm)", "Box direction"],
         [10, 10, 7, 7, 14, 40, 8, 22, 14, 14, 80])
    smap = _sku_map(res)
    r = 2
    for m in res["mixed"]:
        by_name = {c["sku"]: c for c in m["contents"]}
        for st in m["steps"]:
            vals = [m["no"], m["carton"]["name"],
                    st["layer_from"] if st["layer_from"] == st["layer_to"] else f"{st['layer_from']}-{st['layer_to']}",
                    st["step"], "PLACE", st["name"], st["count"], st["grid"],
                    round(st["x"]), round(st["y"]), st["orientation"]]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=r, column=c, value=v)
                cell.border = BORDER
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            r += 1
    ws.freeze_panes = "A2"

    # ---------------------------------------------------------------- Full cartons
    ws = wb.create_sheet("Full Cartons")
    _hdr(ws, 1, ["Carton nos", "Product code", "Product", "Cartons", "Inner boxes / carton", "Units / inner box",
                 "Units / carton", "Carton", "Carton size (in)", "Note"], [14, 14, 44, 9, 12, 10, 12, 10, 20, 70])
    r = 2
    for f in res["full"]:
        no = f"{f['carton_from']} TO {f['carton_to']}" if f["n"] > 1 else str(f["carton_from"])
        vals = [no, f["code"], f["name"], f["n"], f["boxes_per_carton"], f["units_per_inner"], f["units_per_carton"],
                f["carton"]["name"], f["carton_label_in"], f["note"]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1

    # ---------------------------------------------------------------- Warnings
    ws = wb.create_sheet("Warnings")
    _hdr(ws, 1, ["Type", "Message"], [18, 140])
    r = 2
    for w in res.get("warnings", []):
        ws.cell(row=r, column=1, value="Warning")
        ws.cell(row=r, column=2, value=w)
        r += 1
    for w in res.get("excluded", []):
        ws.cell(row=r, column=1, value="Excluded")
        ws.cell(row=r, column=2, value=f"{w['name']}: {w['reason']}")
        r += 1
    for w in res.get("skipped", []):
        ws.cell(row=r, column=1, value="NOT PACKED")
        ws.cell(row=r, column=2, value=f"{w['name']}: {w['reason']}")
        r += 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
