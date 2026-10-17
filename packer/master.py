"""Load the master data file (inner box sizes, boxes per full carton, outer carton sizes)."""
from __future__ import annotations

import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field, asdict
from typing import Optional

import openpyxl

MM_PER_INCH = 25.4


def norm(s) -> str:
    """Aggressive normalisation used for matching: upper-case alphanumerics only."""
    if s is None:
        return ""
    if isinstance(s, float) and s.is_integer():
        s = int(s)
    return re.sub(r"[^A-Z0-9]", "", str(s).upper())


def to_float(v) -> Optional[float]:
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        m = re.search(r"-?\d+(\.\d+)?", str(v))
        return float(m.group()) if m else None


@dataclass
class MasterItem:
    row: int
    code: str            # may be ""
    name: str
    l: Optional[float]   # inner box size in mm (already contains the extra margin)
    w: Optional[float]
    h: Optional[float]
    units_per_inner: Optional[int]
    inner_per_carton: Optional[int]
    carton_key: str      # key of the home outer carton
    carton_label: str

    def issues(self):
        """Human-readable list of what is missing/invalid in this master row."""
        out = []
        if not self.has_dims:
            out.append("inner box size L/W/H is blank or invalid (e.g. -5)")
        if not self.units_per_inner:
            out.append("'ONE INNER BOX Under Patch QTY' (units per inner box) is blank")
        if not self.inner_per_carton:
            out.append("'ONE OUTER BOX UNDER INNER BOX QTY' (inner boxes per full carton) is blank")
        if not self.carton_key:
            out.append("outer carton size (L/W/H in inches) is blank")
        return out

    @property
    def has_dims(self) -> bool:
        return all(v is not None and v > 0 for v in (self.l, self.w, self.h))

    def to_dict(self):
        d = asdict(self)
        d["has_dims"] = self.has_dims
        return d


@dataclass
class CartonSpec:
    key: str
    name: str
    l_in: float
    w_in: float
    h_in: float
    usage: int = 0

    @property
    def l(self):
        return round(self.l_in * MM_PER_INCH, 1)

    @property
    def w(self):
        return round(self.w_in * MM_PER_INCH, 1)

    @property
    def h(self):
        return round(self.h_in * MM_PER_INCH, 1)

    @property
    def label(self):
        f = lambda x: ("%g" % x)
        return f"{f(self.l_in)} x {f(self.w_in)} x {f(self.h_in)}"

    def to_dict(self):
        return {"key": self.key, "name": self.name, "l_in": self.l_in, "w_in": self.w_in,
                "h_in": self.h_in, "l": self.l, "w": self.w, "h": self.h,
                "label": self.label, "usage": self.usage}


@dataclass
class Master:
    items: list = field(default_factory=list)
    cartons: dict = field(default_factory=dict)       # key -> CartonSpec
    by_code: dict = field(default_factory=lambda: defaultdict(list))
    by_name: dict = field(default_factory=lambda: defaultdict(list))

    def index(self):
        self.by_code = defaultdict(list)
        self.by_name = defaultdict(list)
        for it in self.items:
            if it.code:
                self.by_code[norm(it.code)].append(it)
            if it.name:
                self.by_name[norm(it.name)].append(it)


def _carton_key(l, w, h) -> str:
    return f"{l:g}x{w:g}x{h:g}"


def load_master(path: str, sheet: str = "Main data") -> Master:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    ws = wb[sheet] if sheet in wb.sheetnames else wb.worksheets[0]
    master = Master()
    names_by_key = defaultdict(Counter)
    usage = Counter()
    dims_by_key = {}

    for r_idx, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        if not row or len(row) < 14:
            continue
        code, name = row[0], row[1]
        if name is None and code is None:
            continue
        l, w, h = (to_float(row[3]), to_float(row[4]), to_float(row[5]))
        upi = to_float(row[6])
        ipc = to_float(row[7])
        cl, cw, ch = to_float(row[10]), to_float(row[11]), to_float(row[12])
        if name is None and code is None:
            continue
        if (l is not None and l <= 0) or (w is not None and w <= 0) or (h is not None and h <= 0):
            l = w = h = None
        if cl and cw and ch:
            key = _carton_key(cl, cw, ch)
            dims_by_key[key] = (cl, cw, ch)
            label = str(row[13]).strip() if row[13] else key
            names_by_key[key][re.sub(r"\s+", "", label).upper()] += 1
            usage[key] += 1
        else:
            key, label = "", ""
        if not name and not code:
            continue
        master.items.append(MasterItem(
            row=r_idx,
            code=str(int(code)) if isinstance(code, float) and code.is_integer() else (str(code).strip() if code is not None else ""),
            name=str(name).strip() if name else "",
            l=l, w=w, h=h,
            units_per_inner=int(upi) if upi else None,
            inner_per_carton=int(ipc) if ipc and ipc > 0 else None,
            carton_key=key, carton_label=label))

    for key, (cl, cw, ch) in dims_by_key.items():
        best = names_by_key[key].most_common(1)[0][0]
        master.cartons[key] = CartonSpec(key=key, name=best, l_in=cl, w_in=cw, h_in=ch, usage=usage[key])
    # canonical carton name on items
    for it in master.items:
        it.carton_label = master.cartons[it.carton_key].name if it.carton_key in master.cartons else ""
    master.index()
    return master
