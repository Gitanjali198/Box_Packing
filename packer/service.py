"""Glue layer used by the web app and the CLI: master data, matching, manual items, planning."""
from __future__ import annotations

import json
import os
import re
import threading
import time
import uuid
from typing import Optional

from .master import CARTON_REDUCTION_INCH, MM_PER_INCH, Master, MasterItem, load_master, norm
from .matcher import Matcher
from .pi_parser import PIDoc, list_sheets, parse_pi
from .planner import Settings, plan

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(BASE, "data")
MASTER_PATH = os.path.join(DATA, "master.xlsx")
CUSTOM_ITEMS = os.path.join(DATA, "custom_items.json")
CUSTOM_CARTONS = os.path.join(DATA, "custom_cartons.json")
ALIASES = os.path.join(DATA, "aliases.json")
UPLOADS = os.path.join(BASE, "uploads")
os.makedirs(UPLOADS, exist_ok=True)

DEFAULT_MARGIN_MM = 2.0


def _load_json(path, default):
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return default


def _save_json(path, obj):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(obj, f, indent=1, ensure_ascii=False)


class Service:
    def __init__(self):
        self.lock = threading.Lock()
        self.jobs = {}
        self.reload_master()

    # ------------------------------------------------------------------ master / cartons
    def reload_master(self):
        self.master: Master = load_master(MASTER_PATH)
        self.matcher = Matcher(self.master, ALIASES)
        self.custom_items = _load_json(CUSTOM_ITEMS, {})
        self.custom_cartons = _load_json(CUSTOM_CARTONS, {})

    def cartons(self) -> dict:
        out = {}
        for k, c in self.master.cartons.items():
            out[k] = c.to_dict()
        for k, c in self.custom_cartons.items():
            l_in, w_in, h_in = float(c["l_in"]), float(c["w_in"]), float(c["h_in"])
            calc_l_in = max(0.1, round(l_in - CARTON_REDUCTION_INCH, 2))
            calc_w_in = max(0.1, round(w_in - CARTON_REDUCTION_INCH, 2))
            calc_h_in = max(0.1, round(h_in - CARTON_REDUCTION_INCH, 2))
            out[k] = {
                "key": k, "name": c["name"],
                "l_in": l_in, "w_in": w_in, "h_in": h_in,
                "calc_l_in": calc_l_in, "calc_w_in": calc_w_in, "calc_h_in": calc_h_in,
                "l": round(l_in * MM_PER_INCH, 1),
                "w": round(w_in * MM_PER_INCH, 1),
                "h": round(h_in * MM_PER_INCH, 1),
                "outer_l": round(l_in * MM_PER_INCH, 1),
                "outer_w": round(w_in * MM_PER_INCH, 1),
                "outer_h": round(h_in * MM_PER_INCH, 1),
                "calc_l": round(calc_l_in * MM_PER_INCH, 1),
                "calc_w": round(calc_w_in * MM_PER_INCH, 1),
                "calc_h": round(calc_h_in * MM_PER_INCH, 1),
                "label": f"{l_in:g} x {w_in:g} x {h_in:g}",
                "calc_label": f"{calc_l_in:g} x {calc_w_in:g} x {calc_h_in:g}",
                "usage": 0, "custom": True
            }
        return out

    def default_allowed(self):
        cs = self.cartons()
        picks = [k for k, c in cs.items() if c.get("usage", 0) >= 6 or c.get("custom")]
        return picks

    def add_custom_carton(self, name, l_in, w_in, h_in):
        key = f"custom-{norm(name)}-{l_in:g}x{w_in:g}x{h_in:g}"
        self.custom_cartons[key] = {"key": key, "name": name, "l_in": l_in, "w_in": w_in, "h_in": h_in}
        _save_json(CUSTOM_CARTONS, self.custom_cartons)
        return key

    def search_master(self, q, limit=25):
        qn = norm(q)
        res = []
        for it in self.master.items:
            hay = norm(it.name) + "|" + norm(it.code)
            if not qn or qn in hay:
                res.append(it)
        res.sort(key=lambda it: (norm(it.name) != qn, len(it.name), it.row))
        return [self.item_json(it) for it in res[:limit]]

    def item_json(self, it: MasterItem):
        d = it.to_dict()
        d["carton"] = self.master.cartons[it.carton_key].name if it.carton_key in self.master.cartons else ""
        return d

    # ------------------------------------------------------------------ PI parsing & matching
    def pi_sheets(self, path):
        if path.lower().endswith(".xls"):
            import pandas as pd
            names = pd.ExcelFile(path).sheet_names
            out = []
            for n in names:
                try:
                    parse_pi(path, n)
                    out.append({"sheet": n, "is_pi": True, "title": ""})
                except Exception:
                    out.append({"sheet": n, "is_pi": False, "title": ""})
            return out
        return list_sheets(path)

    def _custom_key(self, line):
        return f"{norm(line.code)}|{norm(line.name)}|{line.units_per_box}"

    def parse(self, path, sheet=None):
        doc = parse_pi(path, sheet)
        out = []
        for line in doc.lines:
            ck = self._custom_key(line)
            entry = {"line": line.to_dict(), "status": "missing", "resolution": None, "note": "",
                     "suggest": None, "candidates": [], "issues": []}
            if ck in self.custom_items:
                ci = self.custom_items[ck]
                entry["status"] = "ok"
                entry["resolution"] = {"type": "custom", **ci}
                entry["note"] = "saved manual item"
                entry["issues"].append({"level": "info", "msg": "Using your saved manual box size (not in master sheet) - "
                                        "add it to the master sheet to make it permanent."})
            else:
                m = self.matcher.match(line)
                if m.item is not None:
                    entry["suggest"] = self.item_json(m.item)
                    entry["candidates"] = [self.item_json(c) for c in m.candidates[:6]]
                    entry["note"] = m.note
                    for msg in m.item.issues():
                        entry["issues"].append({"level": "error" if "box size" in msg or "units per inner" in msg else "warn",
                                                "msg": f"Master sheet row {m.item.row} ('{m.item.name}'): {msg}"})
                    if m.note:
                        entry["issues"].append({"level": "info", "msg": m.note + f" (master row {m.item.row})"})
                    if m.status == "fuzzy" and not m.note:
                        entry["issues"].append({"level": "warn", "msg": f"Only an approximate name match to master row {m.item.row} '{m.item.name}' - confirm it is the same product."})
                    if m.item.has_dims:
                        # Auto-resolve using master item box dimensions so item is never ignored/dropped
                        entry["resolution"] = {"type": "master", "row": m.item.row}
                        if m.status in ("exact_code", "exact_name", "manual"):
                            entry["status"] = "ok"
                        else:
                            entry["status"] = "review"
                    else:
                        entry["status"] = "review"
                        entry["note"] = (entry["note"] + " " if entry["note"] else "") + "master has no box dimensions"
                else:
                    entry["note"] = m.note
                    entry["issues"].append({"level": "error", "msg": "Product is NOT in the master sheet "
                                            f"(searched code '{line.code}' and name '{line.name}'). Add it to the master sheet or enter its box size here."})
            out.append(entry)
        return {"sheet": doc.sheet, "title": doc.title, "consignee": doc.consignee, "lines": out,
                "has_highlight": doc.has_highlight}

    def save_custom_item(self, line, spec, remember_master_row=None):
        """spec: {l,w,h (raw mm), per_carton|None, carton_key|None, top_up|False}"""
        ck = self._custom_key(line)
        self.custom_items[ck] = {"code": line["code"], "name": line["name"], "upb": line["units_per_box"],
                                 "l": float(spec["l"]), "w": float(spec["w"]), "h": float(spec["h"]),
                                 "per_carton": int(spec["per_carton"]) if spec.get("per_carton") else None,
                                 "carton_key": spec.get("carton_key") or None,
                                 "top_up": bool(spec.get("top_up")),
                                 "margin": float(spec.get("margin", DEFAULT_MARGIN_MM))}
        _save_json(CUSTOM_ITEMS, self.custom_items)

    def remember_master(self, line, row):
        for it in self.master.items:
            if it.row == row:
                self.matcher.remember(line, it)
                return

    # ------------------------------------------------------------------ build specs & run
    def build_specs(self, lines, settings_adjust=0.0):
        """lines: list of {line: {...}, resolution: {...}|None, skip: bool}"""
        specs = []
        skipped = []
        for e in lines:
            ln = e["line"]
            res = e.get("resolution")
            if e.get("skip"):
                continue          # deliberately excluded (e.g. not yellow-highlighted)
            if not res:
                skipped.append({"idx": ln["idx"], "name": ln["name"], "reason": "ignored - product / box size not available in master data (not included in the calculation)"})
                continue
            try:
                ln["boxes"] = int(ln["boxes"]); ln["units_per_box"] = int(ln["units_per_box"])
            except Exception:
                skipped.append({"idx": ln.get("idx"), "name": ln.get("name", "?"), "reason": "ignored - order quantity / units per box unreadable in the PI"})
                continue
            if res["type"] == "master":
                it = next((x for x in self.master.items if x.row == res["row"]), None)
                if it is None or not it.has_dims:
                    skipped.append({"idx": ln["idx"], "name": ln["name"], "reason": "ignored - master sheet has no box size for this product (not included in the calculation)"})
                    continue
                specs.append({"idx": ln["idx"], "code": ln["code"] or it.code, "name": ln["name"],
                              "boxes": ln["boxes"], "upb": ln["units_per_box"],
                              "l": it.l + settings_adjust, "w": it.w + settings_adjust, "h": it.h + settings_adjust,
                              "per_carton": it.inner_per_carton, "carton_key": it.carton_key,
                              "top_up": it.top_up,
                              "source": f"master row {it.row}"})
            else:
                m = float(res.get("margin", DEFAULT_MARGIN_MM)) + settings_adjust
                specs.append({"idx": ln["idx"], "code": ln["code"], "name": ln["name"],
                              "boxes": ln["boxes"], "upb": ln["units_per_box"],
                              "l": float(res["l"]) + m, "w": float(res["w"]) + m, "h": float(res["h"]) + m,
                              "per_carton": int(res["per_carton"]) if res.get("per_carton") else None,
                              "carton_key": res.get("carton_key") or None,
                              "top_up": bool(res.get("top_up")),
                              "source": "manual item"})
        return specs, skipped

    def start_job(self, lines, settings_dict, margin_adjust=0.0, meta=None):
        jid = uuid.uuid4().hex[:10]
        job = {"id": jid, "status": "running", "progress": 0.0, "result": None, "error": None, "started": time.time()}
        self.jobs[jid] = job
        settings = Settings.from_dict(settings_dict)
        if not settings.allowed_cartons:
            settings.allowed_cartons = self.default_allowed()

        def run():
            try:
                specs, skipped = self.build_specs(lines, margin_adjust)
                def prog(p, k):
                    job["progress"] = round(p, 2)
                res = plan(specs, self.cartons(), settings, progress=prog)
                res["skipped"] = skipped
                res["meta"] = meta or {}
                job["result"] = res
                job["status"] = "done"
                job["progress"] = 1.0
            except Exception as ex:  # noqa
                import traceback
                traceback.print_exc()
                job["status"] = "error"
                job["error"] = str(ex)

        t = threading.Thread(target=run, daemon=True)
        t.start()
        return jid


    # ------------------------------------------------------------------ master maintenance helpers
    HEADERS = ["ITEM CODE", "Top Side Head", "ITEM", "Inner Box Size in mm", "L", "W", "H", "ONE INNER BOX Under Patch QTY",
               "ONE OUTER BOX               UNDER INNER BOX QTY.)", "ONE OUTER Under Patch QTY.",
               "OUTER BOX SIZE In Inch", "L", "W", "H", "NAME OF OUTER"]

    def missing_template(self, entries) -> bytes:
        """Excel sheet (master layout) listing every PI product whose master data is missing/incomplete."""
        import io
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = Workbook(); ws = wb.active; ws.title = "Missing master data"
        ws["A1"] = "Fill the yellow cells, then paste these rows into the 'Main data' sheet of your master file and upload it again."
        ws["A1"].font = Font(bold=True)
        heads = self.HEADERS + ["PROBLEM (delete this column before pasting)", "PI line"]
        for c, h in enumerate(heads, 1):
            cell = ws.cell(row=2, column=c, value=h); cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        yellow = PatternFill("solid", fgColor="FFFF00")
        r = 3
        for e in entries:
            ln = e["line"]; sug = e.get("suggest") or {}
            it = next((x for x in self.master.items if x.row == sug.get("row")), None)
            probs = "; ".join(i["msg"] for i in e.get("issues", []) if i["level"] in ("error", "warn")) or "check"
            vals = [ln["code"] or (it.code if it else ""),
                    "Top Side Head" if (it and it.top_up) else "",
                    it.name if it else ln["name"], "", None, None, None,
                    (it.units_per_inner if it and it.units_per_inner else ln["units_per_box"]),
                    it.inner_per_carton if it else None, None, "", None, None, None, ""]
            if it:
                if it.has_dims:
                    vals[4:7] = [it.l, it.w, it.h]
                if it.carton_key and it.carton_key in self.master.cartons:
                    c = self.master.cartons[it.carton_key]
                    vals[10:15] = [c.label, c.l_in, c.w_in, c.h_in, c.name]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=r, column=c, value=v)
                if v in (None, "") and c != 2:
                    cell.fill = yellow
            ws.cell(row=r, column=16, value=probs)
            ws.cell(row=r, column=17, value=f"{ln['name']} (qty/box {ln['units_per_box']}, {ln['boxes']} boxes)")
            r += 1
        for col, w in zip("ABCDEFGHIJKLMNOPQ", [12, 14, 36, 20, 8, 8, 8, 14, 16, 12, 20, 8, 8, 8, 14, 70, 44]):
            ws.column_dimensions[col].width = w
        buf = io.BytesIO(); wb.save(buf); return buf.getvalue()

    def master_with_custom(self) -> bytes:
        """Copy of the master sheet with the user's saved manual items appended (ready to use as new master)."""
        import io
        import openpyxl
        wb = openpyxl.load_workbook(MASTER_PATH)
        ws = wb["Main data"] if "Main data" in wb.sheetnames else wb.worksheets[0]
        r = ws.max_row + 1
        for ci in self.custom_items.values():
            m = float(ci.get("margin", DEFAULT_MARGIN_MM))
            upi = ci.get("upb")
            ipc = ci.get("per_carton")
            cs = self.cartons().get(ci.get("carton_key") or "")
            top_val = "Top Side Head" if ci.get("top_up") else None
            row = [ci.get("code") or None, top_val, ci.get("name"), f"{ci['l']:g} X {ci['w']:g} X {ci['h']:g}",
                   ci["l"] + m, ci["w"] + m, ci["h"] + m, upi, ipc, (upi * ipc if upi and ipc else None),
                   cs and f"{cs['l_in']:g} x {cs['w_in']:g} x {cs['h_in']:g}", cs and cs["l_in"], cs and cs["w_in"],
                   cs and cs["h_in"], cs and cs["name"]]
            for c, v in enumerate(row, 1):
                ws.cell(row=r, column=c, value=v or None)
            r += 1
        buf = io.BytesIO(); wb.save(buf); return buf.getvalue()
